from pathlib import Path
import re

try:
    import fitz
except Exception:  # pragma: no cover - optional parser dependency.
    fitz = None
try:
    from docx import Document as DocxDocument
except Exception:  # pragma: no cover - optional parser dependency.
    DocxDocument = None
try:
    from openpyxl import load_workbook
except Exception:  # pragma: no cover - optional parser dependency.
    load_workbook = None
try:
    from pypdf import PdfReader
except Exception:  # pragma: no cover - optional parser dependency.
    PdfReader = None


SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx", ".xlsx", ".png", ".jpg", ".jpeg", ".webp"}
NOISE_PATTERNS = [
    r"此为试读",
    r"需要完整\s*PDF\s*请访问",
    r"www\.ertongbook\.com",
    r"ertongbook\.com",
    r"仅供试读",
    r"扫码阅读",
    r"版权所有",
    r"本文档仅供",
    r"#?\s*Page\s+\d+",
    r"第\s*\d+\s*页\s*/\s*共\s*\d+\s*页",
]
NOISE_RE = re.compile("|".join(f"(?:{pattern})" for pattern in NOISE_PATTERNS), re.IGNORECASE)
URL_RE = re.compile(r"https?://\S+|www\.\S+|\b[a-z0-9.-]+\.(?:com|cn|net|org)\b", re.IGNORECASE)
PAGE_MARKER_RE = re.compile(r"^# Page (?P<page>\d+)\n(?P<text>.*?)(?=^# Page \d+\n|\Z)", re.MULTILINE | re.DOTALL)


def _read_text_file(path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "gbk"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="ignore")


def _parse_pdf(path: Path) -> str:
    pages: list[str] = []
    if fitz is not None:
        try:
            with fitz.open(path) as doc:
                for index, page in enumerate(doc, start=1):
                    text = page.get_text("text", sort=True).strip()
                    if text:
                        pages.append(f"# Page {index}\n{text}")
        except Exception:
            pages = []

    if not pages and PdfReader is not None:
        reader = PdfReader(str(path))
        for index, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                pages.append(f"# Page {index}\n{text}")

    result = "\n\n".join(pages).strip()
    if not result:
        raise ValueError("PDF 未提取到可搜索文本。该文件可能是扫描件或图片型 PDF，当前 MVP 暂未内置 OCR。")
    return result


def split_pdf_pages(text: str) -> list[tuple[int, str]]:
    pages = [(int(match.group("page")), match.group("text").strip()) for match in PAGE_MARKER_RE.finditer(text)]
    return pages or [(1, text)]


def clean_extracted_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    cleaned_lines: list[str] = []
    for line in text.splitlines():
        line = NOISE_RE.sub("", line)
        line = URL_RE.sub("", line)
        line = re.sub(r"[ \t]+", " ", line).strip(" ,，:：#")
        if line:
            cleaned_lines.append(line)
    cleaned = "\n".join(cleaned_lines)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    return cleaned.strip()


def chinese_char_count(text: str) -> int:
    return len(re.findall(r"[\u4e00-\u9fff]", text))


def noise_text_ratio(text: str) -> float:
    if not text.strip():
        return 1.0
    noise_chars = sum(len(match.group(0)) for match in NOISE_RE.finditer(text))
    noise_chars += sum(len(match.group(0)) for match in URL_RE.finditer(text))
    page_chars = sum(len(match.group(0)) for match in re.finditer(r"#?\s*Page\s+\d+|第\s*\d+\s*页", text, re.IGNORECASE))
    return min(1.0, (noise_chars + page_chars) / max(len(text), 1))


def is_noise_chunk(text: str, min_chinese_chars: int = 30, max_noise_ratio: float = 0.4) -> bool:
    if NOISE_RE.search(text) or noise_text_ratio(text) > max_noise_ratio:
        return True
    cleaned = clean_extracted_text(text)
    if not cleaned:
        return True
    chinese_count = chinese_char_count(cleaned)
    return chinese_count > 0 and chinese_count < min_chinese_chars


def parse_document(path: Path) -> str:
    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {ext}")

    if ext == ".txt":
        return _read_text_file(path)
    if ext == ".md":
        return _read_text_file(path)
    if ext == ".pdf":
        return _parse_pdf(path)
    if ext == ".docx":
        if DocxDocument is None:
            raise RuntimeError("python-docx is not installed")
        doc = DocxDocument(str(path))
        return "\n".join(paragraph.text for paragraph in doc.paragraphs)
    if ext == ".xlsx":
        if load_workbook is None:
            raise RuntimeError("openpyxl is not installed")
        workbook = load_workbook(str(path), read_only=True, data_only=True)
        rows: list[str] = []
        for sheet in workbook.worksheets:
            rows.append(f"# Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                rows.append("\t".join("" if cell is None else str(cell) for cell in row))
        return "\n".join(rows)

    raise ValueError(f"Unsupported file type: {ext}")
